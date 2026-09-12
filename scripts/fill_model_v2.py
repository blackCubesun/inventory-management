import pandas as pd
import os
from openpyxl import load_workbook

PROCESSED = r'C:\Users\92585\Desktop\Sklad_System\data\processed'
NEW_MODEL = os.path.join(PROCESSED, 'Модель_v2.xlsx')

# Что заливаем
sources = {
    '_Расход': 'Оборотка_год_плоский.xlsx',
    '_Остатки': 'Остатки_плоский.xlsx',
    '_Оборотка': 'Оборотка_месяц_плоский.xlsx',
    '_Резервы': 'Резервы_плоский.xlsx',
    '_ВПути': 'ВПути_плоский.xlsx',
    '_СТО': 'Справочник_СТО.xlsx',
    '_ТочкиЗаказа': 'ТочкиЗаказа.xlsx',
    '_Croston': 'Croston_Прогноз.xlsx',
}

with pd.ExcelWriter(NEW_MODEL, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
    for sheet, filename in sources.items():
        try:
            df = pd.read_excel(os.path.join(PROCESSED, filename))
            df.to_excel(writer, sheet_name=sheet, index=False)
            print(f'  ✓ {sheet}: {df.shape}')
        except Exception as e:
            print(f'  ✗ {sheet}: {e}')

# Скрываем листы
wb = load_workbook(NEW_MODEL)
for sheet in sources.keys():
    if sheet in wb.sheetnames:
        wb[sheet].sheet_state = 'hidden'

wb.save(NEW_MODEL)
print('\nВсе листы загружены и скрыты')